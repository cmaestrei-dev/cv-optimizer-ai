import io
import re

from docx import Document
from docx.enum.text import WD_ALIGN_PARAGRAPH
from docx.shared import Inches, Pt, RGBColor

from models import UserProfile
from services.pdf_generator import clean_markdown_output, strip_emojis

_BOLD_PATTERN = re.compile(r"\*\*(.+?)\*\*")


def _add_inline(paragraph, text: str) -> None:
    # split() con grupo alterna texto normal y **negrita**
    for i, chunk in enumerate(_BOLD_PATTERN.split(text)):
        if chunk:
            paragraph.add_run(chunk).bold = i % 2 == 1


def _setup_styles(doc: Document) -> None:
    normal = doc.styles["Normal"]
    normal.font.name = "Calibri"
    normal.font.size = Pt(10.5)
    for name, size in (("Heading 1", 12), ("Heading 2", 10.5)):
        style = doc.styles[name]
        style.font.name = "Calibri"
        style.font.size = Pt(size)
        style.font.bold = True
        style.font.color.rgb = RGBColor(0x11, 0x11, 0x11)
    for section in doc.sections:
        section.top_margin = section.bottom_margin = Inches(0.6)
        section.left_margin = section.right_margin = Inches(0.7)


def generate_docx(cv_markdown: str, profile: UserProfile) -> bytes:
    """DOCX de una sola columna, sin tablas ni encabezados de página: lo más legible para un ATS."""
    doc = Document()
    _setup_styles(doc)

    if profile.full_name:
        name = doc.add_paragraph()
        name.alignment = WD_ALIGN_PARAGRAPH.CENTER
        run = name.add_run(strip_emojis(profile.full_name.upper()))
        run.bold = True
        run.font.size = Pt(16)
    if profile.contact_line_text:
        contact = doc.add_paragraph(strip_emojis(profile.contact_line_text))
        contact.alignment = WD_ALIGN_PARAGRAPH.CENTER

    for raw_line in strip_emojis(clean_markdown_output(cv_markdown)).splitlines():
        line = raw_line.strip()
        if not line:
            continue
        if line.startswith("### "):
            doc.add_heading(line[4:], level=2)
        elif line.startswith(("## ", "# ")):
            doc.add_heading(line.lstrip("# "), level=1)
        elif line.startswith(("- ", "* ")):
            _add_inline(doc.add_paragraph(style="List Bullet"), line[2:])
        else:
            _add_inline(doc.add_paragraph(), line)

    buffer = io.BytesIO()
    doc.save(buffer)
    return buffer.getvalue()


def build_docx_filename(pdf_filename: str) -> str:
    return pdf_filename.rsplit(".", 1)[0] + ".docx"
